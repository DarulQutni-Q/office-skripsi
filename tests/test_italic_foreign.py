#!/usr/bin/env python3
"""Tests untuk scripts/italic-foreign.py (italic hanya untuk istilah asing).

Usage:
    python3 -m unittest tests.test_italic_foreign -v
"""
import importlib.util
import io
import os
import shutil
import sys
import tempfile
import unittest
import zipfile
from contextlib import redirect_stdout

from docx import Document
from lxml import etree

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_spec = importlib.util.spec_from_file_location(
    "italic_foreign", os.path.join(REPO_ROOT, "scripts", "italic-foreign.py"))
itf = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(itf)

W = itf.W_NS
NS = 'xmlns:w="%s"' % W
q = itf.q


def run_xml(text, rpr=""):
    return '<w:r>%s<w:t xml:space="preserve">%s</w:t></w:r>' % (
        "<w:rPr>%s</w:rPr>" % rpr if rpr else "", text)


def para_xml(*runs, style=None, ppr_extra=""):
    ppr = ""
    if style or ppr_extra:
        ppr = "<w:pPr>%s%s</w:pPr>" % (
            '<w:pStyle w:val="%s"/>' % style if style else "", ppr_extra)
    return "<w:p>%s%s</w:p>" % (ppr, "".join(runs))


def document(*paras):
    return etree.fromstring(
        ('<w:document %s><w:body>%s</w:body></w:document>' % (NS, "".join(paras))).encode("utf-8"))


STYLES = etree.fromstring(("""
<w:styles %s>
  <w:style w:type="paragraph" w:styleId="Normal"><w:name w:val="Normal"/></w:style>
  <w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/>
    <w:pPr><w:outlineLvl w:val="0"/></w:pPr></w:style>
  <w:style w:type="paragraph" w:styleId="GambarCaption"><w:name w:val="Gambar Caption"/></w:style>
  <w:style w:type="paragraph" w:styleId="TOC1"><w:name w:val="toc 1"/></w:style>
  <w:style w:type="paragraph" w:styleId="Kutipan"><w:name w:val="Kutipan"/><w:rPr><w:i/></w:rPr></w:style>
  <w:style w:type="character" w:styleId="Emphasis"><w:name w:val="Emphasis"/><w:rPr><w:i/></w:rPr></w:style>
</w:styles>""" % NS).encode("utf-8"))


def make_ctx(apply=True, strip=False, terms=None, headings=False, bib=False):
    regex = itf.build_regex(terms if terms is not None else itf.load_terms())
    return itf.Ctx(regex, itf.StyleIndex(STYLES), apply, strip, headings, bib)


def runs_of(root, para_index=0):
    """[(teks, miring_langsung)] untuk paragraf ke-N (hanya format langsung)."""
    p = list(root.iter(q("p")))[para_index]
    out = []
    for r in p.iter(q("r")):
        txt = "".join(t.text or "" for t in r.findall(q("t")))
        out.append((txt, itf._flag_state(r.find(q("rPr")), "i")))
    return out


def italic_text(root, para_index=0):
    return [t for t, i in runs_of(root, para_index) if i]


class TestTerms(unittest.TestCase):
    def test_default_regex_boundaries(self):
        rx = itf.build_regex(itf.load_terms())
        self.assertTrue(rx.search("memakai machine learning"))
        self.assertTrue(rx.search("Machine   Learning"))          # case + spasi ganda
        self.assertTrue(rx.search("semua requirements"))           # jamak Inggris
        self.assertIsNone(rx.search("committee"))                  # bukan 'commit'
        self.assertIsNone(rx.search("debugger"))                   # bukan 'debug'
        self.assertTrue(rx.search("di-deploy ke server"))          # imbuhan + hyphen

    def test_exclusion_and_extra_terms(self):
        terms = itf.load_terms(extra=["!software", "state-of-the-art", "# komentar"])
        self.assertNotIn("software", terms)
        self.assertIn("state-of-the-art", terms)
        self.assertNotIn("komentar", terms)

    def test_terms_file(self):
        with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False, encoding="utf-8") as fh:
            fh.write("# istilah saya\nzero trust  # inline\n!website\n")
            path = fh.name
        try:
            terms = itf.load_terms(files=[path])
        finally:
            os.unlink(path)
        self.assertIn("zero trust", terms)
        self.assertNotIn("website", terms)


class TestApply(unittest.TestCase):
    def test_split_single_run_and_keep_formatting(self):
        root = document(para_xml(run_xml(
            "Sistem memakai machine learning untuk klasifikasi.",
            '<w:rFonts w:ascii="Times New Roman"/><w:b/><w:color w:val="FF0000"/><w:sz w:val="24"/>')))
        rep = itf.process_tree(root, make_ctx())
        self.assertEqual(italic_text(root), ["machine learning"])
        self.assertEqual("".join(t for t, _ in runs_of(root)),
                         "Sistem memakai machine learning untuk klasifikasi.")
        self.assertEqual(len(rep.to_italicize), 1)
        for r in root.iter(q("r")):                      # bold/warna/font tetap
            rpr = r.find(q("rPr"))
            self.assertIsNotNone(rpr.find(q("b")))
            self.assertEqual(rpr.find(q("color")).get(q("val")), "FF0000")
            self.assertEqual(rpr.find(q("rFonts")).get(q("ascii")), "Times New Roman")

    def test_rpr_child_order_is_schema_valid(self):
        root = document(para_xml(run_xml(
            "pakai deployment", '<w:rFonts w:ascii="A"/><w:b/><w:color w:val="000000"/><w:sz w:val="24"/><w:lang w:val="id-ID"/>')))
        itf.process_tree(root, make_ctx())
        for r in root.iter(q("r")):
            names = [etree.QName(c).localname for c in r.find(q("rPr"))]
            idx = [itf.RPR_ORDER.index(n) for n in names]
            self.assertEqual(idx, sorted(idx), names)
        self.assertIn("i", names)
        self.assertIn("iCs", names)

    def test_term_fragmented_across_runs(self):
        root = document(para_xml(
            run_xml("Kami pakai machine "), run_xml("learning", "<w:b/>"), run_xml(" dalam riset.")))
        itf.process_tree(root, make_ctx())
        self.assertEqual("".join(italic_text(root)), "machine learning")
        self.assertEqual("".join(t for t, _ in runs_of(root)), "Kami pakai machine learning dalam riset.")

    def test_idempotent(self):
        root = document(para_xml(run_xml("gunakan big data dan dataset")))
        itf.process_tree(root, make_ctx())
        first = etree.tostring(root)
        rep = itf.process_tree(root, make_ctx())
        self.assertEqual(first, etree.tostring(root))
        self.assertEqual(rep.pending, 0)
        self.assertEqual(rep.already, 2)

    def test_already_italic_untouched(self):
        root = document(para_xml(run_xml("machine learning", "<w:i/><w:iCs/>")))
        before = etree.tostring(root)
        itf.process_tree(root, make_ctx())
        self.assertEqual(before, etree.tostring(root))

    def test_explicit_italic_off_is_turned_on(self):
        root = document(para_xml(run_xml("machine learning", '<w:i w:val="0"/>')))
        itf.process_tree(root, make_ctx())
        self.assertEqual(italic_text(root), ["machine learning"])
        self.assertEqual(len(list(root.iter(q("i")))), 1)  # tidak duplikat

    def test_tab_breaks_match(self):
        root = document(para_xml(
            '<w:r><w:t>machine</w:t><w:tab/><w:t>learning</w:t></w:r>'))
        itf.process_tree(root, make_ctx())
        self.assertEqual(italic_text(root), [])

    def test_non_text_children_preserved(self):
        root = document(para_xml(
            '<w:r><w:t xml:space="preserve">pakai machine learning </w:t><w:tab/><w:t>selesai</w:t></w:r>'))
        itf.process_tree(root, make_ctx())
        self.assertEqual(len(list(root.iter(q("tab")))), 1)
        self.assertEqual("".join(t for t, _ in runs_of(root)), "pakai machine learning selesai")

    def test_xml_entities_and_whitespace(self):
        root = document(para_xml(run_xml("AT&amp;T  memakai backend  ")))
        itf.process_tree(root, make_ctx())
        self.assertEqual("".join(t for t, _ in runs_of(root)), "AT&T  memakai backend  ")
        ser = etree.tostring(root).decode()
        self.assertIn("AT&amp;T", ser)
        self.assertNotIn("&amp;amp;", ser)

    def test_trailing_space_preserved_on_split_pieces(self):
        root = document(para_xml(run_xml("lihat dashboard ini")))
        itf.process_tree(root, make_ctx())
        for t in root.iter(q("t")):
            if t.text != t.text.strip():
                self.assertEqual(t.get("{%s}space" % itf.XML_NS), "preserve")

    def test_inside_hyperlink_and_ins(self):
        root = document(
            '<w:p><w:hyperlink><w:r><w:t>lihat dashboard</w:t></w:r></w:hyperlink></w:p>',
            '<w:p><w:ins w:id="1" w:author="a"><w:r><w:t>pakai deploy</w:t></w:r></w:ins></w:p>')
        itf.process_tree(root, make_ctx())
        self.assertEqual(italic_text(root, 0), ["dashboard"])
        self.assertEqual(italic_text(root, 1), ["deploy"])
        self.assertEqual(root.xpath("//w:ins/w:r", namespaces={"w": W}).__len__(), 2)


class TestSkips(unittest.TestCase):
    def test_caption_toc_heading_skipped(self):
        root = document(
            para_xml(run_xml("Gambar 2.1 Arsitektur backend"), style="GambarCaption"),
            para_xml(run_xml("2.1 Backend framework"), style="TOC1"),
            para_xml(run_xml("BAB II Backend"), style="Heading1"),
            para_xml(run_xml("Body backend")))
        rep = itf.process_tree(root, make_ctx())
        self.assertEqual([italic_text(root, i) for i in range(3)], [[], [], []])
        self.assertEqual(italic_text(root, 3), ["backend"])
        self.assertEqual(rep.skipped["caption"], 1)
        self.assertEqual(rep.skipped["toc"], 1)
        self.assertEqual(rep.skipped["heading"], 1)

    def test_heading_included_on_request(self):
        root = document(para_xml(run_xml("Pengujian backend"), style="Heading1"))
        itf.process_tree(root, make_ctx(headings=True))
        self.assertEqual(italic_text(root), ["backend"])

    def test_field_based_detection_without_styles(self):
        root = document(
            '<w:p><w:r><w:t>Tabel </w:t></w:r><w:r><w:instrText> SEQ Tabel \\* ARABIC </w:instrText></w:r>'
            '<w:r><w:t> backend</w:t></w:r></w:p>',
            '<w:p><w:r><w:instrText> PAGEREF _Toc1 \\h </w:instrText></w:r><w:r><w:t>backend 12</w:t></w:r></w:p>')
        rep = itf.process_tree(root, make_ctx())
        self.assertEqual(rep.skipped["caption"], 1)
        self.assertEqual(rep.skipped["toc"], 1)
        self.assertEqual(rep.pending, 0)

    def test_bibliography_skipped_until_lampiran(self):
        root = document(
            para_xml(run_xml("Bab ini membahas backend")),
            para_xml(run_xml("DAFTAR PUSTAKA")),
            para_xml(run_xml("Smith, J. (2020). Backend systems. Journal.")),
            para_xml(run_xml("LAMPIRAN")),
            para_xml(run_xml("Lampiran memakai backend")))
        itf.process_tree(root, make_ctx())
        self.assertEqual(italic_text(root, 0), ["backend"])
        self.assertEqual(italic_text(root, 2), [])
        self.assertEqual(italic_text(root, 4), ["backend"])

    def test_toc_entry_does_not_start_bibliography(self):
        root = document(
            para_xml(run_xml("DAFTAR PUSTAKA"), style="TOC1"),
            para_xml(run_xml("Isi tentang backend")))
        itf.process_tree(root, make_ctx())
        self.assertEqual(italic_text(root, 1), ["backend"])

    def test_bibliography_included_on_request(self):
        root = document(para_xml(run_xml("DAFTAR PUSTAKA")), para_xml(run_xml("Smith. Backend.")))
        itf.process_tree(root, make_ctx(bib=True))
        self.assertEqual(italic_text(root, 1), ["Backend"])

    def test_monospace_run_skipped(self):
        root = document(para_xml(
            run_xml("jalankan ", ""), run_xml("php artisan backend", '<w:rFonts w:ascii="Consolas"/>')))
        itf.process_tree(root, make_ctx())
        self.assertEqual(italic_text(root), [])

    def test_math_runs_ignored(self):
        root = etree.fromstring((
            '<w:document %s xmlns:m="%s"><w:body><w:p><m:oMath><m:r><m:t>backend</m:t></m:r></m:oMath>'
            '</w:p></w:body></w:document>' % (NS, itf.M_NS)).encode())
        itf.process_tree(root, make_ctx())
        self.assertEqual(len(list(root.iter(q("i")))), 0)

    def test_nested_textbox_paragraph_not_double_processed(self):
        root = document(
            '<w:p><w:r><w:pict><w:txbxContent><w:p><w:r><w:t>backend</w:t></w:r></w:p>'
            '</w:txbxContent></w:pict></w:r></w:p>')
        itf.process_tree(root, make_ctx())
        self.assertEqual(len(list(root.iter(q("i")))), 1)


class TestStripOther(unittest.TestCase):
    def test_whole_sentence_italic_keeps_only_foreign(self):
        root = document(para_xml(run_xml(
            "Sistem ini memakai machine learning untuk klasifikasi.", "<w:i/><w:iCs/>")))
        rep = itf.process_tree(root, make_ctx(strip=True))
        self.assertEqual(italic_text(root), ["machine learning"])
        self.assertEqual("".join(t for t, _ in runs_of(root)),
                         "Sistem ini memakai machine learning untuk klasifikasi.")
        self.assertTrue(rep.to_strip)

    def test_fully_non_foreign_italic_run_cleared_in_place(self):
        root = document(para_xml(run_xml("Kalimat biasa", "<w:b/><w:i/><w:iCs/>")))
        itf.process_tree(root, make_ctx(strip=True))
        r = list(root.iter(q("r")))[0]
        self.assertIsNone(r.find(q("rPr")).find(q("i")))
        self.assertIsNotNone(r.find(q("rPr")).find(q("b")))  # bold tetap

    def test_empty_rpr_removed(self):
        root = document(para_xml(run_xml("Kalimat biasa", "<w:i/>")))
        itf.process_tree(root, make_ctx(strip=True))
        self.assertIsNone(list(root.iter(q("r")))[0].find(q("rPr")))

    def test_style_italic_paragraph_gets_explicit_off(self):
        root = document(para_xml(run_xml("Teks kutipan biasa"), style="Kutipan"))
        itf.process_tree(root, make_ctx(strip=True))
        i = list(root.iter(q("i")))[0]
        self.assertEqual(i.get(q("val")), "0")

    def test_style_italic_paragraph_foreign_term_toggle_safe(self):
        # paragraph style italic: istilah asing memang sudah miring -> jangan diubah
        root = document(para_xml(run_xml("pakai backend"), style="Kutipan"))
        rep = itf.process_tree(root, make_ctx())
        self.assertEqual(rep.pending, 0)

    def test_character_style_toggle(self):
        # rStyle Emphasis (italic) + paragraph style italic => saling batal => TIDAK miring
        root = document(para_xml(run_xml("pakai backend", '<w:rStyle w:val="Emphasis"/>'), style="Kutipan"))
        rep = itf.process_tree(root, make_ctx())
        self.assertEqual(rep.pending, 1)
        self.assertEqual(italic_text(root)[0], "backend")

    def test_strip_not_applied_inside_skipped_paragraphs(self):
        root = document(para_xml(run_xml("Judul bab", "<w:i/>"), style="Heading1"))
        itf.process_tree(root, make_ctx(strip=True))
        self.assertEqual(italic_text(root), ["Judul bab"])

    def test_structural_italic_reported_not_changed(self):
        root = document(para_xml(run_xml("Gambar 1 Arsitektur", "<w:i/>"), style="GambarCaption"))
        rep = itf.process_tree(root, make_ctx(strip=True))
        self.assertEqual(len(rep.structural), 1)
        self.assertEqual(italic_text(root), ["Gambar 1 Arsitektur"])


class TestAuditMode(unittest.TestCase):
    def test_audit_does_not_modify(self):
        root = document(para_xml(run_xml("pakai backend")))
        before = etree.tostring(root)
        rep = itf.process_tree(root, make_ctx(apply=False))
        self.assertEqual(before, etree.tostring(root))
        self.assertEqual(len(rep.to_italicize), 1)


class TestEndToEndDocx(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.src = os.path.join(self.tmp, "skripsi.docx")
        doc = Document()
        doc.add_heading("BAB I PENDAHULUAN backend", level=1)
        p = doc.add_paragraph()
        p.add_run("Sistem memakai ")
        p.add_run("machine ")
        r = p.add_run("learning")
        r.bold = True
        p.add_run(" untuk klasifikasi.")
        p2 = doc.add_paragraph()
        r2 = p2.add_run("Seluruh kalimat ini salah miring backend.")
        r2.italic = True
        doc.add_paragraph("Gambar 1 Arsitektur backend", style="Caption")
        doc.add_paragraph("DAFTAR PUSTAKA")
        doc.add_paragraph("Smith (2020). Backend systems.")
        doc.save(self.src)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _cli(self, *args):
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = itf.main(list(args))
        return code, buf.getvalue()

    def test_audit_then_apply_then_clean_audit(self):
        code, out = self._cli("audit", self.src)
        self.assertEqual(code, 1)
        self.assertIn("machine", out)

        dst = os.path.join(self.tmp, "out.docx")
        code, _ = self._cli("apply", self.src, "-o", dst, "--strip-other")
        self.assertEqual(code, 0)

        doc = Document(dst)                                      # tetap terbuka di python-docx
        by_text = {p.text: p for p in doc.paragraphs}
        main = by_text["Sistem memakai machine learning untuk klasifikasi."]
        self.assertEqual("".join(r.text for r in main.runs if r.italic), "machine learning")
        self.assertTrue([r for r in main.runs if r.bold and r.italic])  # 'learning' bold+italic

        wrong = by_text["Seluruh kalimat ini salah miring backend."]
        self.assertEqual("".join(r.text for r in wrong.runs if r.italic), "backend")

        self.assertFalse(any(r.italic for r in by_text["Gambar 1 Arsitektur backend"].runs))
        self.assertFalse(any(r.italic for r in by_text["Smith (2020). Backend systems."].runs))
        head = [p for p in doc.paragraphs if p.text.startswith("BAB I")][0]
        self.assertFalse(any(r.italic for r in head.runs))

        # sudah bersih -> audit exit 0 (dan apply kedua tidak mengubah apa-apa)
        code, out = self._cli("audit", dst, "--strip-other")
        self.assertEqual(code, 0, out)

    def test_original_untouched_and_other_parts_identical(self):
        with open(self.src, "rb") as fh:
            before = fh.read()
        dst = os.path.join(self.tmp, "out.docx")
        self._cli("apply", self.src, "-o", dst)
        with open(self.src, "rb") as fh:
            self.assertEqual(before, fh.read())
        with zipfile.ZipFile(self.src) as a, zipfile.ZipFile(dst) as b:
            self.assertEqual(a.namelist(), b.namelist())
            for n in a.namelist():
                if n != "word/document.xml":
                    self.assertEqual(a.read(n), b.read(n), n)
            etree.fromstring(b.read("word/document.xml"))        # well-formed

    def test_refuses_to_overwrite_input(self):
        code, _ = self._cli("apply", self.src, "-o", self.src)
        self.assertEqual(code, 2)

    def test_dry_run_writes_nothing(self):
        dst = os.path.join(self.tmp, "dry.docx")
        code, _ = self._cli("apply", self.src, "-o", dst, "--dry-run")
        self.assertEqual(code, 0)
        self.assertFalse(os.path.exists(dst))

    def test_unpacked_dir_and_xml_in_place(self):
        folder = os.path.join(self.tmp, "unpacked")
        with zipfile.ZipFile(self.src) as z:
            z.extractall(folder)
        code, _ = self._cli("apply", folder)
        self.assertEqual(code, 0)
        with open(os.path.join(folder, "word", "document.xml"), "rb") as fh:
            xml = fh.read()
        self.assertIn(b"<w:i/>", xml)
        code, _ = self._cli("audit", os.path.join(folder, "word", "document.xml"))
        self.assertEqual(code, 0)

    def test_custom_terms_file(self):
        terms = os.path.join(self.tmp, "t.txt")
        with open(terms, "w", encoding="utf-8") as fh:
            fh.write("klasifikasi\n!machine learning\n")
        dst = os.path.join(self.tmp, "c.docx")
        self._cli("apply", self.src, "-o", dst, "--terms", terms)
        doc = Document(dst)
        main = [p for p in doc.paragraphs if p.text.startswith("Sistem")][0]
        self.assertEqual("".join(r.text for r in main.runs if r.italic), "klasifikasi")


if __name__ == "__main__":
    unittest.main()
