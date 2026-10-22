#!/usr/bin/env python3
"""Tests untuk scripts/validate-docx.py: --type presets (skripsi/makalah).

Usage:
    python3 -m unittest tests.test_validate_docx -v
"""
import contextlib
import importlib.util
import io
import os
import subprocess
import sys
import tempfile
import unittest

from docx import Document

SCRIPTS = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "scripts")
spec = importlib.util.spec_from_file_location(
    "validate_docx", os.path.join(SCRIPTS, "validate-docx.py"))
vd = importlib.util.module_from_spec(spec)
spec.loader.exec_module(vd)

MAKALAH_SECTIONS = ["PENDAHULUAN", "PEMBAHASAN", "PENUTUP", "DAFTAR PUSTAKA"]
SKRIPSI_SECTIONS = ["PENDAHULUAN", "TINJAUAN PUSTAKA", "METODE PENELITIAN",
                    "HASIL DAN PEMBAHASAN", "KESIMPULAN", "DAFTAR PUSTAKA"]


def make_docx(path, headings):
    doc = Document()
    for h in headings:
        doc.add_heading(h, level=1)
        doc.add_paragraph("Isi bagian " + h.lower() + ".")
    doc.save(path)


def quiet_validate(path, sections):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        return vd.validate_docx(path, required_sections=sections)


class TestPresets(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="validate_docx_")
        self.addCleanup(__import__("shutil").rmtree, self.tmp, ignore_errors=True)
        self.makalah = os.path.join(self.tmp, "makalah.docx")
        self.skripsi = os.path.join(self.tmp, "skripsi.docx")
        make_docx(self.makalah, MAKALAH_SECTIONS)
        make_docx(self.skripsi, SKRIPSI_SECTIONS)

    def test_makalah_preset_passes_on_makalah(self):
        self.assertTrue(quiet_validate(self.makalah, vd.DOC_TYPES["makalah"]))

    def test_default_is_skripsi_and_fails_on_makalah(self):
        # backward compat: no sections arg -> skripsi preset -> makalah FAILS.
        # This is the old false-negative, now opt-in via --type.
        self.assertFalse(quiet_validate(self.makalah, None))

    def test_skripsi_preset_passes_on_skripsi(self):
        self.assertTrue(quiet_validate(self.skripsi, vd.DOC_TYPES["skripsi"]))

    def test_makalah_preset_fails_on_skripsi(self):
        # presets are strict: skripsi doc lacks PENUTUP/PEMBAHASAN
        self.assertFalse(quiet_validate(self.skripsi, vd.DOC_TYPES["makalah"]))

    def test_custom_sections_override(self):
        self.assertTrue(quiet_validate(self.makalah, ["PENDAHULUAN", "PENUTUP"]))
        self.assertFalse(quiet_validate(self.makalah, ["PENDAHULUAN", "BAB X"]))

    def test_empty_sections_skips_check(self):
        self.assertTrue(quiet_validate(self.makalah, []))

    def test_zip_and_xml_checks_still_run(self):
        # corrupt file -> still fails at check 1 regardless of sections
        bad = os.path.join(self.tmp, "bad.docx")
        with open(bad, "w") as f:
            f.write("bukan zip")
        self.assertFalse(quiet_validate(bad, []))


class TestResolveSections(unittest.TestCase):
    def test_preset(self):
        self.assertEqual(vd.resolve_sections("skripsi", None), SKRIPSI_SECTIONS)
        self.assertEqual(vd.resolve_sections("makalah", None), MAKALAH_SECTIONS)

    def test_none_skips(self):
        self.assertIsNone(vd.resolve_sections("none", None))

    def test_custom_overrides_preset(self):
        self.assertEqual(vd.resolve_sections("skripsi", "A, B"),
                         ["A", "B"])
        self.assertEqual(vd.resolve_sections("makalah", "PENDAHULUAN"),
                         ["PENDAHULUAN"])

    def test_custom_trims_and_drops_empty(self):
        self.assertEqual(vd.resolve_sections("skripsi", " A ,, B "), ["A", "B"])
        self.assertIsNone(vd.resolve_sections("skripsi", " , "))


class TestCLI(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="validate_docx_cli_")
        self.addCleanup(__import__("shutil").rmtree, self.tmp, ignore_errors=True)
        self.makalah = os.path.join(self.tmp, "makalah.docx")
        make_docx(self.makalah, MAKALAH_SECTIONS)

    def run_cli(self, *args):
        return subprocess.run(
            [sys.executable, os.path.join(SCRIPTS, "validate-docx.py"),
             self.makalah, *args],
            capture_output=True, text=True)

    def test_cli_makalah_exit_0(self):
        self.assertEqual(self.run_cli("--type", "makalah").returncode, 0)

    def test_cli_default_exit_1_on_makalah(self):
        self.assertEqual(self.run_cli().returncode, 1)

    def test_cli_custom_sections_exit_0(self):
        p = self.run_cli("--sections", "PENDAHULUAN, PENUTUP")
        self.assertEqual(p.returncode, 0)

    def test_cli_none_exit_0(self):
        self.assertEqual(self.run_cli("--type", "none").returncode, 0)


if __name__ == "__main__":
    unittest.main()
