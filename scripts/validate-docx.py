#!/usr/bin/env python3
"""
DOCX Validation Tool -- check integrity, structure, and content of a .docx file.

Supports several Indonesian academic document types via --type:

    skripsi  PENDAHULUAN, TINJAUAN PUSTAKA, METODE PENELITIAN,
             HASIL DAN PEMBAHASAN, KESIMPULAN, DAFTAR PUSTAKA
    makalah  PENDAHULUAN, PEMBAHASAN, PENUTUP, DAFTAR PUSTAKA
    none     skip the section-presence check entirely

Or pass your own list:  --sections "PENDAHULUAN, ISI, PENUTUP"

Usage:
    python validate-docx.py revised.docx
    python validate-docx.py makalah.docx --type makalah
    python validate-docx.py bebas.docx --sections "BAB 1, BAB 2"
    python validate-docx.py bebas.docx --type none
"""
import argparse
import sys
import zipfile
from docx import Document
from lxml import etree

# Preset section lists per document type. Matched as case-sensitive
# substrings against word/document.xml (Indonesian academic headings
# are conventionally uppercase).
DOC_TYPES = {
    "skripsi": [
        "PENDAHULUAN", "TINJAUAN PUSTAKA", "METODE PENELITIAN",
        "HASIL DAN PEMBAHASAN", "KESIMPULAN", "DAFTAR PUSTAKA",
    ],
    "makalah": [
        "PENDAHULUAN", "PEMBAHASAN", "PENUTUP", "DAFTAR PUSTAKA",
    ],
}


def resolve_sections(doc_type, sections_arg):
    """Return the section list to check, or None to skip the check."""
    if sections_arg is not None:
        sections = [s.strip() for s in sections_arg.split(",")]
        sections = [s for s in sections if s]
        return sections or None
    if doc_type == "none":
        return None
    return list(DOC_TYPES[doc_type])


def validate_docx(filepath, required_sections=None):
    """Validate a .docx file.

    required_sections: list of section headings that must be present
        (substring match on word/document.xml). None (default) uses the
        "skripsi" preset; pass [] to skip the section check.
    Returns True when every check passes.
    """
    if required_sections is None:
        required_sections = DOC_TYPES["skripsi"]
    print(f"=== VALIDATING {filepath} ===\n")

    passed = 0
    failed = 0

    # 1. ZIP integrity
    print("1. ZIP INTEGRITY")
    try:
        with zipfile.ZipFile(filepath, "r") as z:
            bad = z.testzip()
            if bad:
                print(f"  FAIL: Corrupted entry: {bad}")
                failed += 1
            else:
                print(f"  PASS: ZIP valid ({len(z.namelist())} files)")
                passed += 1
    except Exception as e:
        print(f"  FAIL: ZIP error: {e}")
        failed += 1
        return False

    # 2. Document opens
    print("\n2. DOCUMENT OPENS")
    try:
        doc = Document(filepath)
        print(f"  PASS: Opens successfully")
        print(f"    Paragraphs: {len(doc.paragraphs)}")
        print(f"    Tables: {len(doc.tables)}")
        image_count = sum(1 for r in doc.part.rels.values() if "image" in r.reltype)
        print(f"    Images: {image_count}")
        passed += 1
    except Exception as e:
        print(f"  FAIL: Error: {e}")
        failed += 1
        return False

    # 3. XML validity
    print("\n3. XML VALIDITY")
    xml_errors = []
    try:
        with zipfile.ZipFile(filepath, "r") as z:
            for name in z.namelist():
                if name.endswith(".xml") and "word/" in name:
                    try:
                        etree.fromstring(z.read(name))
                    except Exception as e:
                        xml_errors.append((name, str(e)))
                        print(f"  FAIL: {name}")
        if xml_errors:
            for name, err in xml_errors:
                print(f"    {name}: {err}")
            failed += 1
        else:
            print(f"  PASS: All XML well-formed")
            passed += 1
    except Exception as e:
        print(f"  FAIL: XML check error: {e}")
        failed += 1

    # 4. Content preservation (section presence)
    print("\n4. CONTENT PRESERVATION")
    if not required_sections:
        print("  SKIP: no required sections (use --type or --sections to enable)")
    else:
        try:
            with zipfile.ZipFile(filepath, "r") as z:
                content = z.read("word/document.xml").decode("utf-8")
                missing = []
                for section in required_sections:
                    if section not in content:
                        missing.append(section)
                        print(f"  FAIL: {section} MISSING!")
                if not missing:
                    print(f"  PASS: All required sections present "
                          f"({", ".join(required_sections)})")
                    passed += 1
                else:
                    failed += 1
        except Exception as e:
            print(f"  FAIL: Content check error: {e}")
            failed += 1

    # 5. Summary
    print(f"\n{"="*40}")
    print(f"RESULTS: {passed} passed, {failed} failed")
    if failed == 0:
        print("ALL CHECKS PASSED")
        return True
    else:
        print("SOME CHECKS FAILED")
        return False


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Validate a .docx file: integrity, structure, sections.")
    parser.add_argument("filepath", help="path to the .docx file")
    parser.add_argument("--type", choices=["skripsi", "makalah", "none"],
                        default="skripsi",
                        help="document type preset for the section check "
                             "(default: skripsi)")
    parser.add_argument("--sections",
                        help="custom section list, comma-separated, "
                             "overrides --type")
    args = parser.parse_args(argv)
    sections = resolve_sections(args.type, args.sections)
    success = validate_docx(args.filepath, required_sections=sections or [])
    sys.exit(0 if success else 1)


if __name__ == "__main__":
    main()
